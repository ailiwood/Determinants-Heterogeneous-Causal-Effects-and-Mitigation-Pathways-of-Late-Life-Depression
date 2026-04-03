"""
阶段7.2：面板流失检验（样本选择偏差评估）
—————————————————————————————————
目的：比较流失组（≤2波）与留存组（≥3波）的基线特征，
      评估面板流失是否构成系统性选择偏差。
输入：data/preprocess_outputs/data1_modeling_ready.xlsx
输出：results/06_endogeneity/panel_attrition_baseline_comparison.xlsx
      results/06_endogeneity/panel_attrition_wave_counts.xlsx
      results/06_endogeneity/fig_attrition_*.png
      results/10_tables_for_paper/endogeneity_attrition.xlsx + .docx
"""

import sys
sys.path.insert(0, 'D:/BaiduSyncdisk/lunwen/model0320/scripts')

import numpy as np
import pandas as pd
import os, warnings
warnings.filterwarnings('ignore')
from scipy import stats

# ── 图像配置 ──
import matplotlib
import matplotlib.pyplot as plt
plt.rcParams['font.sans-serif'] = ['SimHei', 'Microsoft YaHei', 'SimSun']
plt.rcParams['axes.unicode_minus'] = False
plt.rcParams['font.size'] = 12
plt.rcParams['figure.dpi'] = 150
plt.rcParams['savefig.dpi'] = 300
plt.rcParams['savefig.bbox'] = 'tight'

# ── 路径 ──
ROOT = 'D:/BaiduSyncdisk/lunwen/model0320'
DATA_PATH = f'{ROOT}/data/preprocess_outputs/data1_modeling_ready.xlsx'
OUT_DIR   = f'{ROOT}/results/06_endogeneity'
TABLE_DIR = f'{ROOT}/results/10_tables_for_paper'
FIG_DIR   = f'{ROOT}/results/11_figures_for_paper'
os.makedirs(OUT_DIR, exist_ok=True)

# ======================================================================
# 1. 读取数据
# ======================================================================
print("=" * 60)
print("7.2  面板流失检验")
print("=" * 60)

print("\n[1/5] 读取数据...")
df = pd.read_excel(DATA_PATH)
print(f"  数据维度: {df.shape}")

# 波次映射
wave_year = {1: 2011, 2: 2013, 3: 2015, 4: 2018, 5: 2020}

# ======================================================================
# 2. 面板结构统计
# ======================================================================
print("\n[2/5] 面板结构统计...")

# 每个ID参与波次数
id_waves = df.groupby('ID')['wave'].agg(['count', 'min', 'max', list])
id_waves.columns = ['n_waves', 'first_wave', 'last_wave', 'waves']

# 分组
id_waves['group'] = id_waves['n_waves'].apply(
    lambda x: '流失组(≤2波)' if x <= 2 else '留存组(≥3波)')

# 波次分布
wave_dist = id_waves['n_waves'].value_counts().sort_index()
print("\n  个体参与波次分布：")
for w, cnt in wave_dist.items():
    pct = cnt / len(id_waves) * 100
    print(f"    {w}波: {cnt} 人 ({pct:.1f}%)")

# 各波次样本量
wave_counts = df.groupby('wave').size().reset_index(name='N')
wave_counts['年份'] = wave_counts['wave'].map(wave_year)
print("\n  各波次样本量：")
for _, row in wave_counts.iterrows():
    print(f"    Wave {int(row['wave'])} ({int(row['年份'])}): {int(row['N'])} 人")

# 保存波次统计
wave_counts_full = pd.DataFrame({
    '参与波次数': wave_dist.index,
    '个体数': wave_dist.values,
    '占比(%)': (wave_dist.values / len(id_waves) * 100).round(1)
})
wave_counts_full.to_excel(f'{OUT_DIR}/panel_attrition_wave_counts.xlsx', index=False)

# ======================================================================
# 3. 基线特征对比（流失组 vs 留存组）
# ======================================================================
print("\n[3/5] 基线特征对比...")

# 提取每个个体的基线观测（首次出现的波次）
df_baseline = df.sort_values(['ID', 'wave']).groupby('ID').first().reset_index()
df_baseline = df_baseline.merge(id_waves[['n_waves', 'group']],
                                 left_on='ID', right_index=True)

print(f"  流失组(≤2波): {(df_baseline['group']=='流失组(≤2波)').sum()} 人")
print(f"  留存组(≥3波): {(df_baseline['group']=='留存组(≥3波)').sum()} 人")

# 要对比的变量
cont_vars = ['cesd10', 'age', 'chronic_disease_count', 'social_activity_index',
             'sleep', 'hhcperc', 'bmi', 'body_discomfort_count', 'family_size']
cat_vars  = ['gender', 'rural2', 'edu', 'ins', 'pension', 'marry']

var_labels = {
    'cesd10': '抑郁得分(CES-D10)',
    'age': '年龄',
    'chronic_disease_count': '慢性病数量',
    'social_activity_index': '社会参与指数',
    'sleep': '睡眠时长(小时)',
    'hhcperc': '人均消费(元)',
    'bmi': '体质指数(BMI)',
    'body_discomfort_count': '身体不适数量',
    'family_size': '家庭规模',
    'gender': '性别(男=1)',
    'rural2': '城镇户口',
    'edu': '受教育程度',
    'ins': '医疗保险',
    'pension': '养老保险',
    'marry': '已婚'
}

g_att = df_baseline[df_baseline['group'] == '流失组(≤2波)']
g_ret = df_baseline[df_baseline['group'] == '留存组(≥3波)']

results = []

# 连续变量：t检验
for var in cont_vars:
    v_att = g_att[var].dropna()
    v_ret = g_ret[var].dropna()
    t_stat, p_val = stats.ttest_ind(v_att, v_ret, equal_var=False)
    d_cohen = (v_att.mean() - v_ret.mean()) / np.sqrt((v_att.std()**2 + v_ret.std()**2) / 2)
    results.append({
        '变量': var_labels.get(var, var),
        '变量名': var,
        '类型': '连续',
        '流失组均值': round(v_att.mean(), 3),
        '流失组SD': round(v_att.std(), 3),
        '留存组均值': round(v_ret.mean(), 3),
        '留存组SD': round(v_ret.std(), 3),
        't/χ²': round(t_stat, 3),
        'p值': round(p_val, 4),
        "Cohen's d": round(d_cohen, 3),
        '显著性': '***' if p_val < 0.01 else ('**' if p_val < 0.05 else ('*' if p_val < 0.1 else ''))
    })

# 分类变量：卡方检验
for var in cat_vars:
    v_att = g_att[var].dropna()
    v_ret = g_ret[var].dropna()

    if var in ['gender', 'rural2', 'ins', 'pension', 'marry']:
        # 二元变量：报告比例
        p_att = v_att.mean()
        p_ret = v_ret.mean()
        contingency = pd.crosstab(
            df_baseline.loc[df_baseline[var].notna(), 'group'],
            df_baseline.loc[df_baseline[var].notna(), var])
        chi2, p_val, _, _ = stats.chi2_contingency(contingency)
        # 效应量 Cramér's V
        n_total = contingency.values.sum()
        cramer_v = np.sqrt(chi2 / n_total)
        results.append({
            '变量': var_labels.get(var, var),
            '变量名': var,
            '类型': '二元',
            '流失组均值': f"{p_att*100:.1f}%",
            '流失组SD': '',
            '留存组均值': f"{p_ret*100:.1f}%",
            '留存组SD': '',
            't/χ²': round(chi2, 3),
            'p值': round(p_val, 4),
            "Cohen's d": round(cramer_v, 3),
            '显著性': '***' if p_val < 0.01 else ('**' if p_val < 0.05 else ('*' if p_val < 0.1 else ''))
        })
    else:
        # 多分类变量（edu）
        p_att = v_att.mean()
        p_ret = v_ret.mean()
        contingency = pd.crosstab(
            df_baseline.loc[df_baseline[var].notna(), 'group'],
            df_baseline.loc[df_baseline[var].notna(), var])
        chi2, p_val, _, _ = stats.chi2_contingency(contingency)
        n_total = contingency.values.sum()
        k = min(contingency.shape) - 1
        cramer_v = np.sqrt(chi2 / (n_total * k)) if k > 0 else 0
        results.append({
            '变量': var_labels.get(var, var),
            '变量名': var,
            '类型': '多分类',
            '流失组均值': round(p_att, 3),
            '流失组SD': round(v_att.std(), 3),
            '留存组均值': round(p_ret, 3),
            '留存组SD': round(v_ret.std(), 3),
            't/χ²': round(chi2, 3),
            'p值': round(p_val, 4),
            "Cohen's d": round(cramer_v, 3),
            '显著性': '***' if p_val < 0.01 else ('**' if p_val < 0.05 else ('*' if p_val < 0.1 else ''))
        })

df_results = pd.DataFrame(results)

# 保存
df_results.to_excel(f'{OUT_DIR}/panel_attrition_baseline_comparison.xlsx', index=False)
df_results.to_excel(f'{TABLE_DIR}/endogeneity_attrition.xlsx', index=False)
print(f"  保存: panel_attrition_baseline_comparison.xlsx")

# ======================================================================
# 4. docx三线表
# ======================================================================
print("\n[4/5] 生成docx三线表...")

from docx import Document
from docx.shared import Pt
from docx.enum.table import WD_TABLE_ALIGNMENT

doc = Document()
doc.add_heading('表：面板流失组与留存组基线特征对比', level=2)
doc.add_paragraph(f'流失组(≤2波) N={len(g_att)}，留存组(≥3波) N={len(g_ret)}')

cols_doc = ['变量', '流失组均值', '流失组SD', '留存组均值', '留存组SD', 't/χ²', 'p值', '显著性']
tbl = doc.add_table(rows=1+len(results), cols=len(cols_doc))
tbl.alignment = WD_TABLE_ALIGNMENT.CENTER

for j, col in enumerate(cols_doc):
    tbl.rows[0].cells[j].text = col
for i, row_data in enumerate(results):
    for j, col in enumerate(cols_doc):
        tbl.rows[i+1].cells[j].text = str(row_data.get(col, ''))

doc.add_paragraph('注：连续变量报告均值(标准差)，采用Welch t检验；'
                  '二元变量报告比例(%)，采用卡方检验。'
                  "Cohen's d为效应量（连续变量）或Cramér's V（分类变量）。"
                  '*** p<0.01, ** p<0.05, * p<0.1。')
doc.save(f'{TABLE_DIR}/endogeneity_attrition.docx')
print(f"  保存: endogeneity_attrition.docx")

# ======================================================================
# 5. 可视化
# ======================================================================
print("\n[5/5] 绘制可视化...")

# 5a. 各波次样本量条形图
fig, ax = plt.subplots(figsize=(8, 5))
waves = wave_counts['wave'].values
years = wave_counts['年份'].values
ns = wave_counts['N'].values
bars = ax.bar(range(len(waves)), ns, color='#2C5F8A', alpha=0.8)
ax.set_xticks(range(len(waves)))
ax.set_xticklabels([f'Wave {int(w)}\n({int(y)})' for w, y in zip(waves, years)])
ax.set_ylabel('样本量')
ax.set_title('各波次有效样本量')
for bar, n in zip(bars, ns):
    ax.text(bar.get_x() + bar.get_width()/2, bar.get_height() + 50,
            str(int(n)), ha='center', fontsize=10)
plt.tight_layout()
fig.savefig(f'{OUT_DIR}/fig_attrition_wave_counts.png')
fig.savefig(f'{FIG_DIR}/endogeneity_wave_counts.png')
plt.close()

# 5b. 个体参与波次数分布
fig, ax = plt.subplots(figsize=(8, 5))
wave_ns = wave_dist.values
wave_labels = [f'{int(w)}波' for w in wave_dist.index]
colors = ['#D4573B' if w <= 2 else '#2C5F8A' for w in wave_dist.index]
bars = ax.bar(range(len(wave_ns)), wave_ns, color=colors, alpha=0.8)
ax.set_xticks(range(len(wave_ns)))
ax.set_xticklabels(wave_labels)
ax.set_ylabel('个体数')
ax.set_title('个体参与波次数分布（红=流失组，蓝=留存组）')
for bar, n in zip(bars, wave_ns):
    ax.text(bar.get_x() + bar.get_width()/2, bar.get_height() + 30,
            str(int(n)), ha='center', fontsize=10)
plt.tight_layout()
fig.savefig(f'{OUT_DIR}/fig_attrition_wave_distribution.png')
fig.savefig(f'{FIG_DIR}/endogeneity_wave_distribution.png')
plt.close()

# 5c. 关键变量分组对比箱线图
key_vars = ['cesd10', 'age', 'chronic_disease_count', 'social_activity_index', 'sleep']
fig, axes = plt.subplots(1, 5, figsize=(20, 5))
for i, var in enumerate(key_vars):
    ax = axes[i]
    data_plot = [g_att[var].dropna().values, g_ret[var].dropna().values]
    bp = ax.boxplot(data_plot, labels=['流失组\n(≤2波)', '留存组\n(≥3波)'],
                    patch_artist=True, widths=0.6)
    bp['boxes'][0].set_facecolor('#D4573B')
    bp['boxes'][0].set_alpha(0.6)
    bp['boxes'][1].set_facecolor('#2C5F8A')
    bp['boxes'][1].set_alpha(0.6)
    ax.set_title(var_labels.get(var, var), fontsize=11)

    # 标注p值
    r = [x for x in results if x['变量名'] == var][0]
    p_text = f"p={r['p值']:.3f}" if r['p值'] >= 0.001 else "p<0.001"
    ax.text(0.5, 0.95, p_text + r['显著性'], transform=ax.transAxes,
            ha='center', va='top', fontsize=9, color='red' if r['p值'] < 0.05 else 'gray')

plt.suptitle('流失组 vs 留存组：关键变量基线对比', fontsize=14, y=1.02)
plt.tight_layout()
fig.savefig(f'{OUT_DIR}/fig_attrition_baseline_comparison.png')
fig.savefig(f'{FIG_DIR}/endogeneity_attrition_baseline.png')
plt.close()

# 5d. 各波次样本进入/退出/留存表
print("\n  样本流动统计：")
wave_list = sorted(df['wave'].unique())
flow_data = []
prev_ids = set()
for w in wave_list:
    cur_ids = set(df[df['wave'] == w]['ID'].unique())
    if len(prev_ids) == 0:
        entered = len(cur_ids)
        exited = 0
        retained = 0
    else:
        entered = len(cur_ids - prev_ids)
        exited = len(prev_ids - cur_ids)
        retained = len(cur_ids & prev_ids)
    flow_data.append({
        'Wave': int(w),
        '年份': wave_year[w],
        '总样本': len(cur_ids),
        '新进入': entered,
        '留存': retained,
        '退出': exited
    })
    prev_ids = cur_ids
    print(f"    Wave {int(w)} ({wave_year[w]}): "
          f"总{len(cur_ids)}, 新进{entered}, 留存{retained}, 退出{exited}")

df_flow = pd.DataFrame(flow_data)
df_flow.to_excel(f'{OUT_DIR}/panel_sample_flow.xlsx', index=False)

print("\n✅ 7.2 面板流失检验完成！")
print(f"\n关键发现：")
for r in results:
    if r['p值'] < 0.05:
        cohen_d = r["Cohen's d"]
        print(f"  ⚠ {r['变量']}: 流失组={r['流失组均值']}, 留存组={r['留存组均值']}, "
              f"p={r['p值']}{r['显著性']}, d={cohen_d}")

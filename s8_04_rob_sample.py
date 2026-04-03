"""
阶段8.4：样本替换稳健性检验 + 8.5汇总表
—————————————————————————————————
(a) 仅保留≥3波个体
(b) 删除年龄>90的极端高龄
(c) 汇总所有稳健性检验结果
"""

import sys
sys.path.insert(0, 'D:/BaiduSyncdisk/lunwen/model0320/scripts')

import numpy as np
import pandas as pd
import json, os, warnings
warnings.filterwarnings('ignore')

from sklearn.preprocessing import LabelEncoder
from xgboost import XGBRegressor, XGBClassifier
import doubleml as dml
from doubleml import DoubleMLData, DoubleMLPLR, DoubleMLIRM

ROOT = 'D:/BaiduSyncdisk/lunwen/model0320'
DATA_LAG = f'{ROOT}/results/03_causal_sample/causal_sample_lagged.xlsx'
DATA_ORIG = f'{ROOT}/data/preprocess_outputs/data1_modeling_ready.xlsx'
OUT_DIR  = f'{ROOT}/results/07_robustness'
TABLE_DIR = f'{ROOT}/results/10_tables_for_paper'
os.makedirs(OUT_DIR, exist_ok=True)
os.makedirs(TABLE_DIR, exist_ok=True)

SEED = 42
N_FOLDS = 5
N_REP = 5

def get_ml_reg():
    return XGBRegressor(n_estimators=500, max_depth=5, learning_rate=0.1,
                        random_state=SEED, n_jobs=-1, verbosity=0)
def get_ml_cls():
    return XGBClassifier(n_estimators=500, max_depth=5, learning_rate=0.1,
                         random_state=SEED, n_jobs=-1, verbosity=0,
                         use_label_encoder=False, eval_metric='logloss')

def stars(p):
    if isinstance(p, str): return ''
    return '***' if p<0.01 else ('**' if p<0.05 else ('*' if p<0.1 else ''))

print("=" * 60)
print("8.4  样本替换稳健性检验")
print("=" * 60)

# 读取因果样本
df_full = pd.read_excel(DATA_LAG)
le = LabelEncoder()
df_full['province_enc'] = le.fit_transform(df_full['province'].astype(str))
print(f"全因果样本: {df_full.shape}")

# 读取原始面板（用于确定≥3波个体）
df_orig = pd.read_excel(DATA_ORIG)
id_wave_counts = df_orig.groupby('ID')['wave'].nunique()
ids_3plus = set(id_wave_counts[id_wave_counts >= 3].index)

W_COMMON = ['age', 'gender', 'edu', 'marry', 'rural2', 'province_enc',
            'hhcperc_log', 'ins', 'pension', 'retire', 'bmi', 'smoken',
            'drinkev', 'family_size', 'hchild', 'lag_cesd10', 'wave']
W_T1 = W_COMMON + ['chronic_disease_count', 'sleep', 'fcamt_log']
W_T2 = W_COMMON + ['sleep', 'social_activity_index', 'fcamt_log']
W_T3 = W_COMMON + ['chronic_disease_count', 'social_activity_index', 'fcamt_log']

results_sample = []

def run_dml_on_subset(df_sub, sub_label):
    """对子样本运行三个treatment的DML"""
    Y = df_sub['cesd10_t'].values
    sub_results = []

    # T1: PLR
    print(f"\n  T1 ({sub_label})")
    d1 = DoubleMLData.from_arrays(y=Y, d=df_sub['social_activity_index'].values, x=df_sub[W_T1].values)
    m1 = DoubleMLPLR(d1, ml_l=get_ml_reg(), ml_m=get_ml_reg(), n_folds=N_FOLDS, n_rep=N_REP, score='partialling out')
    m1.fit()
    print(f"    θ={m1.coef[0]:.4f} (SE={m1.se[0]:.4f}, p={m1.pval[0]:.4f})")
    sub_results.append({'检验项': f'T1 {sub_label}', '模型': 'PLR',
                        'ATE/θ': round(m1.coef[0], 4), 'SE': round(m1.se[0], 4),
                        'p值': round(m1.pval[0], 4), '显著性': stars(m1.pval[0])})

    # T2: IRM
    print(f"  T2 ({sub_label})")
    d2 = DoubleMLData.from_arrays(y=Y, d=df_sub['T2_chronic'].values, x=df_sub[W_T2].values)
    m2 = DoubleMLIRM(d2, ml_g=get_ml_reg(), ml_m=get_ml_cls(), n_folds=N_FOLDS, n_rep=N_REP,
                     score='ATE', trimming_threshold=0.01)
    m2.fit()
    print(f"    ATE={m2.coef[0]:.4f} (SE={m2.se[0]:.4f}, p={m2.pval[0]:.4f})")
    sub_results.append({'检验项': f'T2 {sub_label}', '模型': 'IRM',
                        'ATE/θ': round(m2.coef[0], 4), 'SE': round(m2.se[0], 4),
                        'p值': round(m2.pval[0], 4), '显著性': stars(m2.pval[0])})

    # T3: IRM
    print(f"  T3 ({sub_label})")
    d3 = DoubleMLData.from_arrays(y=Y, d=df_sub['T3_sleep'].values, x=df_sub[W_T3].values)
    m3 = DoubleMLIRM(d3, ml_g=get_ml_reg(), ml_m=get_ml_cls(), n_folds=N_FOLDS, n_rep=N_REP,
                     score='ATE', trimming_threshold=0.01)
    m3.fit()
    print(f"    ATE={m3.coef[0]:.4f} (SE={m3.se[0]:.4f}, p={m3.pval[0]:.4f})")
    sub_results.append({'检验项': f'T3 {sub_label}', '模型': 'IRM',
                        'ATE/θ': round(m3.coef[0], 4), 'SE': round(m3.se[0], 4),
                        'p值': round(m3.pval[0], 4), '显著性': stars(m3.pval[0])})

    return sub_results

# === (a) 仅≥3波个体 ===
df_3w = df_full[df_full['ID'].isin(ids_3plus)].copy()
print(f"\n--- 子样本a: ≥3波个体 ---")
print(f"  N={df_3w.shape[0]}, IDs={df_3w['ID'].nunique()}")
res_a = run_dml_on_subset(df_3w, '≥3波')
results_sample.extend(res_a)

df_a = pd.DataFrame(res_a)
df_a.to_excel(f'{OUT_DIR}/rob_sample_3waves.xlsx', index=False)

# === (b) 排除>90岁 ===
df_no90 = df_full[df_full['age'] <= 90].copy()
print(f"\n--- 子样本b: 排除>90岁 ---")
print(f"  N={df_no90.shape[0]} (排除{df_full.shape[0]-df_no90.shape[0]}条)")
res_b = run_dml_on_subset(df_no90, '排除>90岁')
results_sample.extend(res_b)

df_b = pd.DataFrame(res_b)
df_b.to_excel(f'{OUT_DIR}/rob_sample_no90.xlsx', index=False)

# ======================================================================
# 8.5 汇总所有稳健性检验
# ======================================================================
print("\n" + "=" * 60)
print("8.5  稳健性检验汇总")
print("=" * 60)

# 读取所有子检验结果
dfs = []
for fname in ['rob_treatment_definition.xlsx', 'rob_outcome_binary.xlsx',
              'rob_learner_all.xlsx', 'rob_sample_3waves.xlsx', 'rob_sample_no90.xlsx']:
    fpath = f'{OUT_DIR}/{fname}'
    if os.path.exists(fpath):
        df_tmp = pd.read_excel(fpath)
        df_tmp['来源'] = fname.replace('.xlsx', '').replace('rob_', '')
        dfs.append(df_tmp)
    else:
        print(f"  警告: {fname} 不存在")

# 读取主模型结果
with open(f'{ROOT}/results/04_dml/dml_results.json', 'r') as f:
    main_dml = json.load(f)

main_rows = []
for key, label in [('T1_xgb', 'T1主模型(DML)'), ('T2_xgb', 'T2主模型(DML)'), ('T3_xgb', 'T3主模型(DML)')]:
    r = main_dml[key]
    if 'theta' in r:
        est, se, p = r['theta'], r['se'], r['pvalue']
    else:
        est, se, p = r['ATE'], r['ATE_se'], r['ATE_pvalue']
    main_rows.append({'检验项': label, '模型': 'DML-XGB(主)',
                      'ATE/θ': round(est, 4), 'SE': round(se, 4),
                      'p值': round(p, 4), '显著性': stars(p), '来源': 'main'})

df_main = pd.DataFrame(main_rows)
dfs.insert(0, df_main)

df_summary = pd.concat(dfs, ignore_index=True)
df_summary.to_excel(f'{OUT_DIR}/robustness_summary.xlsx', index=False)
df_summary.to_excel(f'{TABLE_DIR}/robustness_summary.xlsx', index=False)

# 生成docx
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT

doc = Document()
doc.add_heading('表：稳健性检验汇总', level=2)

cols = ['检验项', '模型', 'ATE/θ', 'SE', 'p值', '显著性']
tbl = doc.add_table(rows=1+len(df_summary), cols=len(cols))
tbl.alignment = WD_TABLE_ALIGNMENT.CENTER

for j, col in enumerate(cols):
    tbl.rows[0].cells[j].text = col
for i, (_, row) in enumerate(df_summary.iterrows()):
    for j, col in enumerate(cols):
        tbl.rows[i+1].cells[j].text = str(row.get(col, ''))

doc.add_paragraph('注：*** p<0.01, ** p<0.05, * p<0.1。DML均采用5折交叉拟合×5次重复。'
                  'PLR用于连续treatment，IRM用于二元treatment。'
                  'CRE为Mundlak相关随机效应模型。LPM-FE为线性概率模型固定效应。')
doc.save(f'{TABLE_DIR}/robustness_summary.docx')

print(f"\n保存: robustness_summary.xlsx + .docx")
print("\n===== 稳健性检验汇总表 =====")
print(df_summary[cols].to_string(index=False))
print("\n✅ 8.4 + 8.5 样本替换与汇总完成！")
